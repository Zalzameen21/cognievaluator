import io
import datetime
from django.utils import timezone
from docx import Document
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet

def generate_student_excel_report(student, attendance_records, eval_stats, evaluations, tasks, batch_comparison=None):
    wb = Workbook()
    
    # 1. Summary Sheet
    ws_summary = wb.active
    ws_summary.title = "Student Summary"
    
    header_font = Font(bold=True, size=14)
    label_font = Font(bold=True)
    
    ws_summary['A1'] = f"Performance Report: {student.name}"
    ws_summary['A1'].font = header_font
    
    ws_summary['A3'] = "Batch:"
    ws_summary['B3'] = student.batch.name
    ws_summary['A4'] = "Roll Number:"
    ws_summary['B4'] = student.roll_number
    ws_summary['A5'] = "Overall Score:"
    ws_summary['B5'] = f"{student.average_score}%"
    
    ws_summary['A7'] = "Attendance Report"
    ws_summary['A7'].font = header_font
    
    total_att = attendance_records.count()
    present_att = attendance_records.filter(status__in=['PRESENT', 'LATE']).count()
    pct = round((present_att / total_att * 100) if total_att else 0, 1)
    
    ws_summary['A9'] = "Total Days:"
    ws_summary['B9'] = total_att
    ws_summary['A10'] = "Days Present:"
    ws_summary['B10'] = present_att
    ws_summary['A11'] = "Attendance %:"
    ws_summary['B11'] = f"{pct}%"
    
    if batch_comparison:
        ws_summary['C11'] = f"(Batch Avg: {batch_comparison['avg_attendance']}%)"
    
    ws_summary['A13'] = "Task Report"
    ws_summary['A13'].font = header_font
    
    total_tasks = tasks.count()
    completed_tasks = tasks.filter(status='COMPLETED').count()
    task_pct = round((completed_tasks / total_tasks * 100) if total_tasks else 0, 1)
    
    ws_summary['A15'] = "Total Tasks Given:"
    ws_summary['B15'] = total_tasks
    ws_summary['A16'] = "Tasks Completed:"
    ws_summary['B16'] = completed_tasks
    ws_summary['A17'] = "Completion %:"
    ws_summary['B17'] = f"{task_pct}%"
    
    if batch_comparison:
        ws_summary['C17'] = f"(Batch Avg: {batch_comparison['avg_task_completion']}%)"
        
    ws_summary['A19'] = "Overall Progress / Performance Trend"
    ws_summary['A19'].font = header_font
    
    row = 21
    evals_asc = list(evaluations)[::-1]
    for idx, ev in enumerate(evals_asc):
        ws_summary[f'A{row}'] = f"Evaluation {idx+1} ({ev.date}):"
        ws_summary[f'B{row}'] = f"{ev.total_score}%"
        row += 1
        
    row += 2
    ws_summary[f'A{row}'] = "Recent Evaluator Remarks"
    ws_summary[f'A{row}'].font = header_font
    row += 2
    for ev in evaluations[:5]:
        ws_summary[f'A{row}'] = f"Date: {ev.date}"
        ws_summary[f'A{row}'].font = label_font
        ws_summary[f'B{row}'] = f"Score: {ev.total_score}% | Evaluator: {ev.evaluator.get_full_name() or ev.evaluator.username}"
        
        ws_summary[f'A{row+1}'] = "Strongest Areas:"
        ws_summary[f'B{row+1}'] = ev.strongest_areas
        ws_summary[f'A{row+2}'] = "Areas for Improvement:"
        ws_summary[f'B{row+2}'] = ev.areas_for_improvement
        ws_summary[f'A{row+3}'] = "General Remarks:"
        ws_summary[f'B{row+3}'] = ev.general_remarks
        row += 5
        
    for col in ws_summary.columns:
        ws_summary.column_dimensions[col[0].column_letter].width = 25

    # 2. Tasks Sheet
    ws_tasks = wb.create_sheet(title="Specific Task Details")
    task_headers = ["Task Title", "Day Given", "Date Given", "Due Date", "Status", "Completed At", "Evaluator Remarks"]
    ws_tasks.append(task_headers)
    
    for c in range(1, len(task_headers)+1):
        cell = ws_tasks.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
        
    for task in tasks:
        ws_tasks.append([
            task.title,
            task.assigned_date.strftime('%A') if task.assigned_date else '',
            task.assigned_date.strftime('%Y-%m-%d') if task.assigned_date else '',
            task.due_date.strftime('%Y-%m-%d') if task.due_date else '',
            task.status,
            task.completed_at.strftime('%Y-%m-%d %H:%M') if task.completed_at else '',
            task.evaluator_remarks
        ])
        
    for col in ws_tasks.columns:
        ws_tasks.column_dimensions[col[0].column_letter].width = 20

    # 3. Attendance Sheet
    ws_att = wb.create_sheet(title="Attendance Dates")
    ws_att.append(["Date", "Status", "Marked By"])
    for c in range(1, 4):
        ws_att.cell(row=1, column=c).font = Font(bold=True)
    
    for att in attendance_records:
        marked_by = att.marked_by.get_full_name() or att.marked_by.username if att.marked_by else "System"
        ws_att.append([att.date.strftime('%Y-%m-%d'), att.status, marked_by])
        
    ws_att.column_dimensions['A'].width = 15
    ws_att.column_dimensions['B'].width = 15
    ws_att.column_dimensions['C'].width = 25

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_student_docx_report(student, attendance_records, eval_stats, evaluations, tasks, batch_comparison=None):
    doc = Document()
    doc.add_heading(f"Performance Report: {student.name}", 0)
    
    doc.add_paragraph(f"Batch: {student.batch.name}")
    doc.add_paragraph(f"Roll Number: {student.roll_number}")
    doc.add_paragraph(f"Overall Score: {student.average_score}%")
    
    # Attendance
    doc.add_heading('Attendance Report', level=1)
    total_att = attendance_records.count()
    present_att = attendance_records.filter(status__in=['PRESENT', 'LATE']).count()
    pct = round((present_att / total_att * 100) if total_att else 0, 1)
    
    doc.add_paragraph(f"Total Days: {total_att}")
    doc.add_paragraph(f"Days Present: {present_att}")
    doc.add_paragraph(f"Attendance %: {pct}%")
    
    if batch_comparison:
        doc.add_paragraph(f"Batch Average Attendance: {batch_comparison['avg_attendance']}%")
    
    doc.add_heading('Detailed Attendance', level=2)
    table_att = doc.add_table(rows=1, cols=3)
    table_att.style = 'Table Grid'
    hdr_att = table_att.rows[0].cells
    hdr_att[0].text = 'Date'
    hdr_att[1].text = 'Status'
    hdr_att[2].text = 'Marked By'
    
    for att in attendance_records:
        row = table_att.add_row().cells
        marked_by = att.marked_by.get_full_name() or att.marked_by.username if att.marked_by else "System"
        row[0].text = att.date.strftime('%Y-%m-%d')
        row[1].text = att.status
        row[2].text = marked_by
    
    # Tasks Summary
    doc.add_heading('Task Report Summary', level=1)
    total_tasks = tasks.count()
    completed_tasks = tasks.filter(status='COMPLETED').count()
    task_pct = round((completed_tasks / total_tasks * 100) if total_tasks else 0, 1)
    
    doc.add_paragraph(f"Total Tasks Given: {total_tasks}")
    doc.add_paragraph(f"Tasks Completed: {completed_tasks}")
    doc.add_paragraph(f"Completion %: {task_pct}%")
    
    if batch_comparison:
        doc.add_paragraph(f"Batch Average Task Completion: {batch_comparison['avg_task_completion']}%")
    
    # Specific Tasks Table
    doc.add_heading('Specific Task Details', level=2)
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Task'
    hdr_cells[1].text = 'Day'
    hdr_cells[2].text = 'Given Date'
    hdr_cells[3].text = 'Status'
    hdr_cells[4].text = 'Completed Date'
    
    for task in tasks:
        row_cells = table.add_row().cells
        row_cells[0].text = task.title
        row_cells[1].text = task.assigned_date.strftime('%A') if task.assigned_date else '-'
        row_cells[2].text = task.assigned_date.strftime('%Y-%m-%d') if task.assigned_date else '-'
        row_cells[3].text = task.status
        row_cells[4].text = task.completed_at.strftime('%Y-%m-%d') if task.completed_at else '-'
        
    # Performance Trend
    doc.add_heading('Overall Progress / Performance Trend', level=1)
    evals_asc = list(evaluations)[::-1]
    for idx, ev in enumerate(evals_asc):
        doc.add_paragraph(f"Evaluation {idx+1} ({ev.date}) → {ev.total_score}%")
        
    # Evaluations
    doc.add_heading('Evaluator Insights', level=1)
    for ev in evaluations[:3]:
        p = doc.add_paragraph()
        p.add_run(f"Date: {ev.date} | Score: {ev.total_score}%\n").bold = True
        if ev.strongest_areas:
            p.add_run("Strongest Areas: ").bold = True
            p.add_run(f"{ev.strongest_areas}\n")
        if ev.areas_for_improvement:
            p.add_run("Areas for Improvement: ").bold = True
            p.add_run(f"{ev.areas_for_improvement}\n")
        if ev.general_remarks:
            p.add_run("Remarks: ").bold = True
            p.add_run(f"{ev.general_remarks}\n")
            
    output = io.BytesIO()
    doc.save(output)
    output.seek(0)
    return output


def generate_student_pdf_report(student, attendance_records, eval_stats, evaluations, tasks, batch_comparison=None):
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter)
    styles = getSampleStyleSheet()
    elements = []
    
    # Title
    elements.append(Paragraph(f"Performance Report: {student.name}", styles['Title']))
    elements.append(Paragraph(f"Batch: {student.batch.name} | Roll Number: {student.roll_number} | Overall Score: {student.average_score}%", styles['Normal']))
    elements.append(Spacer(1, 12))
    
    # Attendance
    elements.append(Paragraph("Attendance Report", styles['Heading2']))
    total_att = attendance_records.count()
    present_att = attendance_records.filter(status__in=['PRESENT', 'LATE']).count()
    pct = round((present_att / total_att * 100) if total_att else 0, 1)
    
    batch_att_str = f" ({batch_comparison['avg_attendance']}% Batch Avg)" if batch_comparison else ""
    
    att_data = [
        ["Total Days", "Days Present", "Attendance %"],
        [str(total_att), str(present_att), f"{pct}%{batch_att_str}"]
    ]
    t_att = Table(att_data, colWidths=[100, 100, 150])
    t_att.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.grey),
        ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,0), 12),
        ('GRID', (0,0), (-1,-1), 1, colors.black)
    ]))
    elements.append(t_att)
    elements.append(Spacer(1, 12))
    
    elements.append(Paragraph("Detailed Attendance", styles['Heading2']))
    att_detail_data = [["Date", "Status", "Marked By"]]
    for att in attendance_records:
        marked_by = att.marked_by.get_full_name() or att.marked_by.username if att.marked_by else "System"
        att_detail_data.append([att.date.strftime('%Y-%m-%d'), att.status, marked_by])
        
    t_att_detail = Table(att_detail_data, colWidths=[100, 100, 150])
    t_att_detail.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.lightgrey),
        ('TEXTCOLOR', (0,0), (-1,0), colors.black),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,0), 8),
        ('GRID', (0,0), (-1,-1), 1, colors.black),
    ]))
    elements.append(t_att_detail)
    elements.append(Spacer(1, 24))
    
    # Tasks Summary
    elements.append(Paragraph("Task Report Summary", styles['Heading2']))
    total_tasks = tasks.count()
    completed_tasks = tasks.filter(status='COMPLETED').count()
    task_pct = round((completed_tasks / total_tasks * 100) if total_tasks else 0, 1)
    
    batch_task_str = f" ({batch_comparison['avg_task_completion']}% Batch Avg)" if batch_comparison else ""
    
    task_data = [
        ["Total Tasks Given", "Tasks Completed", "Completion %"],
        [str(total_tasks), str(completed_tasks), f"{task_pct}%{batch_task_str}"]
    ]
    t_task = Table(task_data, colWidths=[120, 120, 150])
    t_task.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.grey),
        ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,0), 12),
        ('GRID', (0,0), (-1,-1), 1, colors.black)
    ]))
    elements.append(t_task)
    elements.append(Spacer(1, 24))
    
    # Task Details Table
    elements.append(Paragraph("Specific Task Details", styles['Heading2']))
    task_table_data = [["Task Title", "Day", "Given Date", "Status", "Completed Date"]]
    for task in tasks:
        task_table_data.append([
            Paragraph(task.title, styles['Normal']),
            task.assigned_date.strftime('%A') if task.assigned_date else '-',
            task.assigned_date.strftime('%Y-%m-%d') if task.assigned_date else '-',
            task.status,
            task.completed_at.strftime('%Y-%m-%d') if task.completed_at else '-'
        ])
        
    t_task_details = Table(task_table_data, colWidths=[150, 70, 70, 70, 90])
    t_task_details.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.lightgrey),
        ('TEXTCOLOR', (0,0), (-1,0), colors.black),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,0), 8),
        ('GRID', (0,0), (-1,-1), 1, colors.black),
        ('VALIGN',(0,0),(-1,-1),'TOP'),
    ]))
    elements.append(t_task_details)
    elements.append(Spacer(1, 24))
    
    # Performance Trend
    elements.append(Paragraph("Overall Progress / Performance Trend", styles['Heading2']))
    evals_asc = list(evaluations)[::-1]
    for idx, ev in enumerate(evals_asc):
        elements.append(Paragraph(f"<b>Evaluation {idx+1} ({ev.date}) &rarr; {ev.total_score}%</b>", styles['Normal']))
    elements.append(Spacer(1, 16))
    
    # Evaluator Remarks
    elements.append(Paragraph("Evaluator Insights & Remarks", styles['Heading2']))
    for ev in evaluations[:3]:
        elements.append(Paragraph(f"<b>Date: {ev.date} (Score: {ev.total_score}%)</b>", styles['Normal']))
        if ev.strongest_areas:
            elements.append(Paragraph(f"<b>Strongest Areas:</b> {ev.strongest_areas}", styles['Normal']))
        if ev.areas_for_improvement:
            elements.append(Paragraph(f"<b>Improvement Needed:</b> {ev.areas_for_improvement}", styles['Normal']))
        if ev.general_remarks:
            elements.append(Paragraph(f"<b>Remarks:</b> {ev.general_remarks}", styles['Normal']))
        elements.append(Spacer(1, 10))
        
    doc.build(elements)
    output.seek(0)
    return output
