import csv
import io
import json
from datetime import datetime
from django.shortcuts import render, get_object_or_404
from django.http import HttpResponse, HttpResponseForbidden
from django.db.models import Avg, Count, Q
from django.template.loader import get_template

from .models import (
    Location, Batch, Student, Attendance, StudentTask,
    DailyEvaluation, Project, BatchScheduleException
)
from .decorators import (
    evaluator_required, get_user_batches, can_user_access_batch
)

@evaluator_required
def documents_hub_view(request):
    user = request.user
    accessible_batches = get_user_batches(user).select_related('location').order_by('-created_at')
    
    locations = {}
    for batch in accessible_batches:
        if batch.location not in locations:
            locations[batch.location] = []
        locations[batch.location].append(batch)
    
    context = {
        'locations_batches': locations,
    }
    return render(request, 'evaluator_app/documents/documents_hub.html', context)


@evaluator_required
def documents_batch_preview_view(request, batch_id):
    batch = get_object_or_404(Batch, id=batch_id)
    if not can_user_access_batch(request.user, batch):
        return HttpResponseForbidden("You do not have permission to access this batch.")
    
    # Aggregations for preview
    students = Student.objects.filter(batch=batch)
    total_students = students.count()
    active_students = students.filter(status='ACTIVE').count()
    
    tasks_assigned = StudentTask.objects.filter(student__batch=batch).count()
    evaluations_completed = DailyEvaluation.objects.filter(student__batch=batch).count()
    projects = Project.objects.filter(batch=batch).count()
    
    context = {
        'batch': batch,
        'total_students': total_students,
        'active_students': active_students,
        'tasks_assigned': tasks_assigned,
        'evaluations_completed': evaluations_completed,
        'projects': projects,
    }
    return render(request, 'evaluator_app/documents/documents_batch_preview.html', context)

@evaluator_required
def documents_download_view(request, batch_id):
    batch = get_object_or_404(Batch, id=batch_id)
    if not can_user_access_batch(request.user, batch):
        return HttpResponseForbidden("You do not have permission to access this batch.")
        
    if request.method == 'POST':
        export_format = request.POST.get('format', 'pdf')
        
        if export_format == 'pdf':
            return _generate_pdf_report(request, batch)
        elif export_format == 'excel':
            return _generate_excel_report(request, batch)
        elif export_format == 'csv':
            return _generate_csv_report(request, batch)
            
    return HttpResponse("Invalid Request", status=400)

def _generate_pdf_report(request, batch):
    try:
        from xhtml2pdf import pisa
    except ImportError:
        return HttpResponse("PDF generation library not installed. Please install xhtml2pdf.", status=500)
        
    # Gather all necessary data
    students = Student.objects.filter(batch=batch).order_by('roll_number')
    
    # Pre-calculate data for each student
    for student in students:
        # Attendance
        total_classes = Attendance.objects.filter(student=student).count()
        present = Attendance.objects.filter(student=student, status='PRESENT').count()
        student.total_classes = total_classes
        student.present = present
        student.absent = Attendance.objects.filter(student=student, status='ABSENT').count()
        student.late = Attendance.objects.filter(student=student, status='LATE').count()
        student.attendance_pct = round((present / total_classes * 100), 1) if total_classes > 0 else 0
        
        # Tasks
        tasks = StudentTask.objects.filter(student=student)
        student.tasks_assigned = tasks.count()
        student.tasks_completed = tasks.filter(status__in=['COMPLETED', 'LATE']).count()
        student.tasks_pending = tasks.filter(status__in=['PENDING', 'OVERDUE']).count()
        student.tasks_needs_revision = tasks.filter(evaluation_status='NEEDS_REVISION').count()
        student.task_completion_pct = round((student.tasks_completed / student.tasks_assigned * 100), 1) if student.tasks_assigned > 0 else 0
        
        # Evaluations
        evals = DailyEvaluation.objects.filter(student=student).order_by('-date')
        student.eval_count = evals.count()
        student.latest_eval = evals.first()
        student.avg_score = round(evals.aggregate(Avg('total_score'))['total_score__avg'] or 0, 1)
        
        # Projects
        student.student_projects = Project.objects.filter(students=student)
    
    template_path = 'evaluator_app/documents/document_pdf_template.html'
    context = {'batch': batch, 'students': students}
    
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="batch_report_{batch.code}.pdf"'
    
    template = get_template(template_path)
    html = template.render(context)
    
    pisa_status = pisa.CreatePDF(
       html, dest=response)
       
    if pisa_status.err:
       return HttpResponse('We had some errors <pre>' + html + '</pre>')
    return response

def _generate_excel_report(request, batch):
    try:
        import openpyxl
    except ImportError:
        return HttpResponse("Excel library not installed. Please install openpyxl.", status=500)
        
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="batch_report_{batch.code}.xlsx"'
    
    wb = openpyxl.Workbook()
    
    # 1. Batch Summary
    ws_summary = wb.active
    ws_summary.title = "Batch Summary"
    ws_summary.append(["Batch Name", batch.name])
    ws_summary.append(["Course Code", batch.code])
    ws_summary.append(["Location", batch.location.name])
    ws_summary.append(["Start Date", str(batch.start_date)])
    ws_summary.append(["End Date", str(batch.end_date) if batch.end_date else "Ongoing"])
    
    students = Student.objects.filter(batch=batch).order_by('roll_number')
    
    # 2. Students
    ws_students = wb.create_sheet("Students")
    ws_students.append(["Roll No", "Name", "Email", "Phone", "Status", "Join Date"])
    for s in students:
        ws_students.append([s.roll_number, s.name, s.email, s.phone, s.get_status_display(), str(s.join_date)])
        
    # 3. Attendance
    ws_attendance = wb.create_sheet("Attendance")
    ws_attendance.append(["Roll No", "Name", "Total Classes", "Present", "Absent", "Late", "Attendance %"])
    for s in students:
        total = Attendance.objects.filter(student=s).count()
        present = Attendance.objects.filter(student=s, status='PRESENT').count()
        absent = Attendance.objects.filter(student=s, status='ABSENT').count()
        late = Attendance.objects.filter(student=s, status='LATE').count()
        pct = round((present / total * 100), 1) if total > 0 else 0
        ws_attendance.append([s.roll_number, s.name, total, present, absent, late, pct])
        
    # 4. Tasks
    ws_tasks = wb.create_sheet("Tasks")
    ws_tasks.append(["Roll No", "Name", "Assigned", "Completed", "Pending", "Needs Revision", "Completion %"])
    for s in students:
        tasks = StudentTask.objects.filter(student=s)
        assigned = tasks.count()
        completed = tasks.filter(status__in=['COMPLETED', 'LATE']).count()
        pending = tasks.filter(status__in=['PENDING', 'OVERDUE']).count()
        revision = tasks.filter(evaluation_status='NEEDS_REVISION').count()
        pct = round((completed / assigned * 100), 1) if assigned > 0 else 0
        ws_tasks.append([s.roll_number, s.name, assigned, completed, pending, revision, pct])
        
    # 5. Evaluations
    ws_evals = wb.create_sheet("Evaluations")
    ws_evals.append(["Roll No", "Name", "Evaluations Count", "Average Score", "Latest Grade"])
    for s in students:
        evals = DailyEvaluation.objects.filter(student=s).order_by('-date')
        count = evals.count()
        avg_score = evals.aggregate(Avg('total_score'))['total_score__avg']
        avg_score = round(avg_score, 1) if avg_score else 0
        latest = evals.first()
        latest_grade = latest.grade if latest else "N/A"
        ws_evals.append([s.roll_number, s.name, count, avg_score, latest_grade])
        
    # 6. Projects
    ws_projects = wb.create_sheet("Projects")
    ws_projects.append(["Roll No", "Name", "Project Title", "Status", "Progress %"])
    for s in students:
        projs = Project.objects.filter(students=s)
        for p in projs:
            ws_projects.append([s.roll_number, s.name, p.title, p.get_status_display(), p.progress_percentage])
            
    wb.save(response)
    return response

def _generate_csv_report(request, batch):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = f'attachment; filename="batch_students_{batch.code}.csv"'
    
    writer = csv.writer(response)
    writer.writerow(['Roll Number', 'Name', 'Phone', 'Email', 'Status'])
    
    students = Student.objects.filter(batch=batch).order_by('roll_number')
    for student in students:
        writer.writerow([student.roll_number, student.name, student.phone, student.email, student.status])
        
    return response
